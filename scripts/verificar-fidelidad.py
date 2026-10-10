"""Compara el texto de las fichas, recetas y preparaciones con el de las
fuentes: cada dato debería aparecer (casi) literal en la página o
diapositiva que cita. Lista lo que no coincide, para revisar a mano.

Uso:  python3 scripts/verificar-fidelidad.py
Requiere: pip install pymupdf python-pptx pyyaml
"""
import glob
import re
import sys
import unicodedata
from pathlib import Path

import pymupdf
import yaml
from pptx import Presentation

RAIZ = Path(__file__).resolve().parent.parent
DOCS = {
    "guia-digestiva": RAIZ / "fuentes/guia_plantas_medicinales.pdf",
    "guia-inmunidad": RAIZ / "fuentes/guia-inmunidad-inteligente-2.pdf",
    "taller-digestivo": RAIZ / "fuentes/taller_plantas_medicinales.pptx",
    "taller-estres": RAIZ / "fuentes/Taller_18-4.pptx",
    "guia-alergia": RAIZ / "fuentes/guia-plantas-alergia.pdf",
}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKC", s).lower()
    s = s.replace("’", "'").replace("’", "'").replace("–", "-").replace("—", "-")
    s = s.replace("*", "")
    s = re.sub(r"[«»\"“”👉⚠️🌿🌙☕🧠➕🌸💜🌱~]", " ", s)
    # Se compara sin puntuación: solo importan las palabras y los números.
    s = re.sub(r"[^\w%/'()+\-]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def textos_pdf(ruta):
    d = pymupdf.open(ruta)
    return {i + 1: norm(p.get_text()) for i, p in enumerate(d)}


def recorrer(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from recorrer(sh.shapes)
        else:
            yield sh


def textos_pptx(ruta):
    pres = Presentation(ruta)
    out = {}
    for i, s in enumerate(pres.slides):
        partes = [sh.text_frame.text for sh in recorrer(s.shapes) if sh.has_text_frame]
        out[i + 1] = norm(" ".join(partes))
    return out


FUENTE = {k: (textos_pdf(v) if v.suffix == ".pdf" else textos_pptx(v)) for k, v in DOCS.items()}


def paginas(ref):
    doc, pags = ref.split("/")
    a, _, b = pags.partition("-")
    return doc, range(int(a), int(b or a) + 1)


def fragmentos(texto):
    """Oraciones o trozos (se cortan en . : ; y en «/»). Los textos cortos van enteros."""
    crudos = re.split(r"(?<=[.:;])\s+|\s+/\s+|\s+·\s+", unicodedata.normalize("NFKC", texto))
    trozos = [norm(x) for x in crudos]
    trozos = [x for x in trozos if len(x.split()) >= 3]
    return trozos or [norm(texto)]


def revisar(texto, refs, donde):
    fuente = " ".join(FUENTE[d][p] for r in refs for d, ps in [paginas(r)] for p in ps)
    faltan = [f for f in fragmentos(texto) if f not in fuente]
    return faltan


problemas = 0
revisados = 0


def chequear(texto, refs, donde):
    global problemas, revisados
    if not refs:
        return
    revisados += 1
    faltan = revisar(texto, refs if isinstance(refs, list) else [refs], donde)
    if faltan:
        problemas += 1
        print(f"· {donde} [{', '.join(refs if isinstance(refs, list) else [refs])}]")
        for f in faltan:
            print(f"    no literal: «{f[:140]}»")


for ruta in sorted(glob.glob(str(RAIZ / "src/content/plantas/*.yaml"))):
    d = yaml.safe_load(open(ruta))
    n = Path(ruta).stem
    for campo in ["usos", "contraindicaciones", "efectosAdversos", "interacciones", "formaDeUso", "comoActua", "evidencia"]:
        for i, it in enumerate(d.get(campo, [])):
            chequear(it["texto"], it["de"], f"plantas/{n} {campo}[{i}]")
    for i, c in enumerate(d.get("cuidados", [])):
        chequear(c["situacion"] + ". " + c["texto"], c["de"], f"plantas/{n} cuidados[{i}]")
    for i, r in enumerate(d.get("recuadros", [])):
        chequear(r["texto"], r["de"], f"plantas/{n} recuadros[{i}]")
    for campo in ["compuestoActivo", "composicion"]:
        if d.get(campo):
            chequear(d[campo]["texto"], d[campo]["de"], f"plantas/{n} {campo}")

for ruta in sorted(glob.glob(str(RAIZ / "src/content/recetas/*.yaml"))):
    d = yaml.safe_load(open(ruta))
    n = Path(ruta).stem
    refs = d["de"] if isinstance(d["de"], list) else [d["de"]]
    for i, p in enumerate(d.get("pasos", [])):
        chequear(p["texto"], refs, f"recetas/{n} pasos[{i}]")
    for campo in ["recomendadoPara", "formaDeUso", "conservacion", "porQue", "precauciones", "notas"]:
        for i, it in enumerate(d.get(campo, [])):
            chequear(it["texto"], it["de"], f"recetas/{n} {campo}[{i}]")

for ruta in sorted(glob.glob(str(RAIZ / "src/content/preparaciones/*.yaml"))):
    d = yaml.safe_load(open(ruta))
    n = Path(ruta).stem
    for s in d["secciones"]:
        for i, it in enumerate(s.get("items", [])):
            chequear(it["texto"], it["de"], f"preparaciones/{n} «{s['titulo']}»[{i}]")

# Glosario y avisos de los usos (src/data).
for e in yaml.safe_load(open(RAIZ / "src/data/glosario.yaml")):
    if e.get("de"):
        chequear(e["definicion"], e["de"], f"glosario/{e['id']}")
for u in yaml.safe_load(open(RAIZ / "src/data/usos.yaml")):
    for i, a in enumerate(u.get("avisos", [])):
        chequear(a["texto"], a["de"], f"usos/{u['id']} avisos[{i}]")

print(f"\n{revisados} datos revisados; {problemas} con fragmentos que no aparecen literales en la fuente citada.")

# Temas (Markdown): cada tramo de texto se compara con la fuente del [[fuente: …]] que lo cierra.
# Se cuentan aparte porque los temas reordenan más el texto (tablas, listas, títulos propios).
def plano_md(t):
    # Las notas de edición (data-edicion) no salen de la fuente: no se comparan.
    t = re.sub(r"<aside[^>]*data-edicion[^>]*>.*?</aside>", " ", t, flags=re.S)
    t = re.sub(r"</?(aside|div)[^>]*>", " ", t)
    t = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", t)
    t = re.sub(r"^#+ .*$", " ", t, flags=re.M)
    t = re.sub(r"^\|?[-| :]+\|?$", " ", t, flags=re.M)
    t = t.replace("|", ". ")
    t = re.sub(r"^\s*[-*] ", "", t, flags=re.M)
    return t

revisados_t = problemas_t = 0
for ruta in sorted(glob.glob(str(RAIZ / "src/content/temas/*.md"))):
    n = Path(ruta).stem
    cuerpo = open(ruta).read().split("---", 2)[2]
    tramos = re.split(r"\[\[fuente:\s*([^\]]+)\]\]", cuerpo)
    for i in range(1, len(tramos), 2):
        refs = [r.strip() for r in tramos[i].split(",")]
        texto = plano_md(tramos[i - 1])
        faltan = [f for f in fragmentos(texto) if f not in " ".join(FUENTE[d][p] for r in refs for d, ps in [paginas(r)] for p in ps)]
        revisados_t += 1
        if faltan:
            problemas_t += 1
            print(f"· temas/{n} [{', '.join(refs)}]")
            for f in faltan:
                print(f"    no literal: «{f[:140]}»")
print(f"\nTemas: {revisados_t} tramos revisados; {problemas_t} con fragmentos que no aparecen literales.")
print("Los que figuran arriba no son necesariamente errores (puede ser una cita recortada o una frase de")
print("enlace), pero conviene mirarlos contra el documento original.")
