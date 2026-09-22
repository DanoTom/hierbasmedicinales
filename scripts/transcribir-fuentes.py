"""Genera transcripciones de texto de los documentos de /fuentes.

Uso:  python3 scripts/transcribir-fuentes.py
Requiere: pip install pymupdf python-pptx

Las transcripciones son una ayuda de lectura. La fuente de verdad son
siempre los archivos originales (PDF y PPTX) de /fuentes.
"""
import re
from pathlib import Path

import pymupdf
from pptx import Presentation

RAIZ = Path(__file__).resolve().parent.parent
FUENTES = RAIZ / "fuentes"
SALIDA = FUENTES / "texto"

# Encabezados y pies que se repiten en cada página y no aportan contenido.
RUIDO_PDF = {
    "GUÍA INMUNIDAD INTELIGENTE",
    "mburucuyá · hierbas",
    "Material de divulgación · no reemplaza la consulta profesional",
}


def limpiar(linea: str) -> str:
    return re.sub(r"\s+", " ", linea).strip()


def lineas_de_pagina(pagina):
    """Devuelve las líneas de una página con su posición y tamaño de letra."""
    lineas = []
    for bloque in pagina.get_text("dict")["blocks"]:
        for linea in bloque.get("lines", []):
            texto = "".join(s["text"] for s in linea["spans"])
            texto = limpiar(texto)
            if not texto:
                continue
            tam = max(s["size"] for s in linea["spans"])
            negrita = all("Bold" in s["font"] or s["flags"] & 16 for s in linea["spans"] if s["text"].strip())
            x0, y0, x1, y1 = linea["bbox"]
            nueva = {"t": texto, "x0": x0, "x1": x1, "y0": y0, "y1": y1, "tam": round(tam, 1), "neg": negrita}
            ant = lineas[-1] if lineas else None
            # Texto justificado: algunas líneas vienen partidas palabra por palabra.
            if ant and abs(ant["y0"] - y0) < 2 and 0 <= x0 - ant["x1"] < 30 and abs(ant["tam"] - nueva["tam"]) < 0.5:
                ant["t"] += " " + texto
                ant["x1"] = x1
            else:
                lineas.append(nueva)
    for l in lineas:
        # Íconos de la tipografía de símbolos (área de uso privado de Unicode).
        l["t"] = limpiar(re.sub(r"[\ue000-\uf8ff]", "", l["t"]))
    return [l for l in lineas if l["t"]]


def pdf_a_md(pdf: Path, titulo: str, nota: str) -> str:
    doc = pymupdf.open(pdf)
    partes = [f"# {titulo}\n", f"> Transcripción de `fuentes/{pdf.name}`. {nota}\n"]
    for n, pagina in enumerate(doc, start=1):
        partes.append(f"\n---\n\n## Página {n} del PDF\n")
        parrafos = []
        previa = None
        vineta = False
        for l in lineas_de_pagina(pagina):
            if l["t"] == "bullet":
                # El PDF dibuja las viñetas con una tipografía de símbolos.
                vineta = True
                continue
            if l["t"] in RUIDO_PDF or re.fullmatch(r"\d{1,2}", l["t"]):
                continue
            if l["t"].startswith("bullet"):
                vineta, l["t"] = True, l["t"][len("bullet"):].strip()
            continua = (
                not vineta
                and previa is not None
                and abs(l["tam"] - previa["tam"]) < 0.5
                and l["neg"] == previa["neg"]
                and 0 <= l["y0"] - previa["y1"] < previa["tam"] * 0.6
                and abs(l["x0"] - previa["x0"]) < 40
            )
            if continua:
                parrafos[-1] += " " + l["t"]
            else:
                prefijo = "• " if vineta else ""
                parrafos.append(prefijo + ("**" if l["neg"] else "") + l["t"])
            vineta = False
            previa = l
        for par in parrafos:
            if par.startswith("**") or par.startswith("• **"):
                par = par + "**"
            partes.append(par + "\n")
    return "\n".join(partes)


def recorrer(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # grupo
            yield from recorrer(sh.shapes)
        else:
            yield sh


def pptx_a_md(pptx: Path, titulo: str, nota: str) -> str:
    pres = Presentation(pptx)
    partes = [f"# {titulo}\n", f"> Transcripción de `fuentes/{pptx.name}`. {nota}\n"]
    for n, diapo in enumerate(pres.slides, start=1):
        partes.append(f"\n---\n\n## Diapositiva {n}\n")
        for sh in recorrer(diapo.shapes):
            if sh.shape_type == 13:
                partes.append("*[imagen]*\n")
            if not sh.has_text_frame:
                continue
            lineas = [limpiar(p.text) for p in sh.text_frame.paragraphs]
            lineas = [l for l in lineas if l]
            if lineas:
                partes.append("\n".join(lineas) + "\n")
        if diapo.has_notes_slide:
            notas = diapo.notes_slide.notes_text_frame.text.strip()
            if notas:
                partes.append(f"**Notas del orador:** {notas}\n")
    return "\n".join(partes)


def main():
    SALIDA.mkdir(exist_ok=True)
    trabajos = [
        ("guia_plantas_medicinales.pdf", "Guía práctica: Plantas medicinales para el sistema digestivo", pdf_a_md,
         "Texto extraído del PDF (tiene texto digital, no hizo falta OCR). Las tablas quedan como bloques de texto seguidos."),
        ("guia-inmunidad-inteligente-2.pdf", "Guía de acompañamiento del taller: Inmunidad inteligente", pdf_a_md,
         "Texto extraído del PDF (tiene texto digital, no hizo falta OCR). Se omitieron el encabezado, el pie ('Material de divulgación · no reemplaza la consulta profesional') y el número de página que se repiten en cada hoja. Las tablas quedan como bloques de texto seguidos."),
        ("taller_plantas_medicinales.pptx", "Presentación del taller: Plantas medicinales para el sistema digestivo", pptx_a_md,
         "Texto de cada diapositiva, en el orden en que aparece en el archivo. Las tablas de las diapositivas 6 a 11 están armadas con cuadros de texto sueltos: cada fila se lee como Planta / Nombre científico / Compuesto activo / Indicación principal / Forma de uso."),
        ("Taller_18-4.pptx", "Presentación del taller: Plantas medicinales para el estrés y el descanso", pptx_a_md,
         "Texto de cada diapositiva. Las imágenes son ilustraciones decorativas y láminas botánicas (melisa, lavanda, valeriana); no contienen texto adicional."),
    ]
    for archivo, titulo, fn, nota in trabajos:
        md = fn(FUENTES / archivo, titulo, nota)
        destino = SALIDA / (Path(archivo).stem + ".md")
        destino.write_text(md, encoding="utf-8")
        print(f"{destino.relative_to(RAIZ)}: {len(md)} caracteres")


if __name__ == "__main__":
    main()
